import io
from datetime import date
from decimal import Decimal as D

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import Company, RevenueShare
from app.services import revenue_shares
from app.services.company_statements import build_statements


def setup(client):
    """Two apps; an investor gets 10% of App A after App Store fees, a consultant 20% of reported revenue."""
    cid = client.post("/api/companies", json={"name": "Apps LLC"}).json()["id"]
    client.post(f"/api/companies/{cid}/members", json={"name": "Jared", "equity_percent": 100})
    app_a = client.post(f"/api/companies/{cid}/products", json={"name": "App A"}).json()["id"]
    app_b = client.post(f"/api/companies/{cid}/products", json={"name": "App B"}).json()["id"]

    investor = client.post(f"/api/companies/{cid}/revenue-shares", json={
        "payee": "Investor", "product_id": app_a, "rate_percent": 10,
        "deduct_categories": ["App Store fees", " app store FEES ", ""], "start_date": "2025-01-01",
    })
    assert investor.status_code == 201, investor.text
    consultant = client.post(f"/api/companies/{cid}/revenue-shares", json={
        "payee": "Consultant", "basis": "reported", "product_id": app_b, "rate_percent": 20,
        "deduct_categories": ["ignored"], "start_date": "2025-02-01", "end_date": "2025-11-30",
        "category": "Commissions",
    })
    assert consultant.status_code == 201, consultant.text

    def add(**body):
        r = client.post(f"/api/companies/{cid}/entries", json={"description": "x", **body})
        assert r.status_code == 201, r.text
        return r.json()

    add(kind="income", date="2025-01-15", amount=1000, product_id=app_a, category="Sales")
    add(kind="expense", date="2025-01-31", amount=150, product_id=app_a, category="App Store fees")
    add(kind="income", date="2025-02-15", amount=2000.55, product_id=app_a, category="Sales")
    add(kind="income", date="2025-02-15", amount=500, product_id=app_b, category="Sales")
    add(kind="income", date="2025-02-20", amount=300, category="Consulting")  # no product: not App A's
    add(kind="expense", date="2025-03-01", amount=40, product_id=app_b, category="Hosting")
    add(kind="expense", date="2025-02-10", amount=85, product_id=app_a, category="Revenue share",
        revenue_share_id=investor.json()["id"])
    return cid, app_a, app_b, investor.json(), consultant.json(), add


def test_products(client):
    cid, app_a, app_b, *_ , add = setup(client)
    assert client.post(f"/api/companies/{cid}/products", json={"name": "App A"}).status_code == 409
    assert [p["name"] for p in client.get(f"/api/companies/{cid}").json()["products"]] == ["App A", "App B"]

    # A contribution never has a product; another company's product isn't allowed.
    member = client.get(f"/api/companies/{cid}").json()["members"][0]["id"]
    c = add(kind="contribution", date="2025-01-01", amount=10, member_id=member, product_id=app_a)
    assert c["product_id"] is None
    other = client.post("/api/companies", json={"name": "Other"}).json()["id"]
    other_product = client.post(f"/api/companies/{other}/products", json={"name": "X"}).json()["id"]
    bad = client.post(f"/api/companies/{cid}/entries", json={
        "kind": "income", "date": "2025-01-01", "description": "x", "amount": 1, "product_id": other_product})
    assert bad.status_code == 404

    s = client.get(f"/api/companies/{cid}/summary", params={"year": 2025}).json()
    assert s["cost_of_revenue"] == 235  # App Store fees 150 + revenue share 85
    rows = {r["name"]: r for r in s["products"]}
    assert list(rows) == ["App A", "App B", "No product"]
    assert (rows["App A"]["income"], rows["App A"]["cost_of_revenue"], rows["App A"]["net"]) == (3000.55, 235, 2765.55)
    assert (rows["App B"]["other_expenses"], rows["No product"]["income"]) == (40, 300)

    # Products in use can't be deleted; unused ones can.
    used = client.delete(f"/api/companies/{cid}/products/{app_a}")
    assert used.status_code == 409 and "4 entries and 1 revenue share" in used.json()["detail"]
    spare = client.post(f"/api/companies/{cid}/products", json={"name": "Spare"}).json()["id"]
    assert client.delete(f"/api/companies/{cid}/products/{spare}").status_code == 204


def test_income_basis(client):
    cid, app_a, _, investor, _, add = setup(client)
    assert investor["deduct_categories"] == ["App Store fees"]
    months = client.get(f"/api/companies/{cid}/revenue-shares/{investor['id']}/months", params={"year": 2025}).json()
    assert len(months) == 12
    jan, feb, mar = months[:3]
    assert (jan["revenue"], jan["deductions"], jan["base"], jan["earned"]) == (1000, 150, 850, 85)
    assert (feb["revenue"], feb["earned"], feb["paid"], feb["balance"]) == (2000.55, 200.06, 85, 200.06)  # rounds half up
    assert (mar["earned"], mar["balance"]) == (0, 200.06)

    with SessionLocal() as db:
        t = revenue_shares.totals(db, db.get(RevenueShare, investor["id"]), 2025, today=date(2026, 1, 5))
    assert (t.earned_year, t.paid_year, t.balance_due) == (D("285.06"), 85, D("200.06"))

    # A month whose fees exceed its income earns nothing rather than going negative.
    add(kind="expense", date="2025-04-30", amount=99, product_id=app_a, category="App Store fees")
    apr = client.get(f"/api/companies/{cid}/revenue-shares/{investor['id']}/months", params={"year": 2025}).json()[3]
    assert (apr["base"], apr["earned"]) == (0, 0)

    shares = client.get(f"/api/companies/{cid}/revenue-shares", params={"year": 2025}).json()
    assert [s["payee"] for s in shares] == ["Consultant", "Investor"]
    assert shares[1]["balance_due"] == 200.06 and shares[1]["paid_total"] == 85


def test_reported_basis(client):
    cid, _, _, investor, consultant, _ = setup(client)
    sid = consultant["id"]
    assert consultant["deduct_categories"] == []
    assert consultant["terms"] == "20% of eligible revenue, reported monthly, starting February 1, 2025 through November 30, 2025"
    assert consultant["missing_months"][:2] == ["2025-02", "2025-03"] and consultant["missing_months"][-1] == "2025-11"

    url = f"/api/companies/{cid}/revenue-shares/{sid}/months"
    assert client.put(url, json={"months": [{"month": "2025-02", "revenue": 1000}, {"month": "2025-03", "revenue": 0},
                                            {"month": "2025-04", "revenue": 10}]}).status_code == 204
    assert client.put(url, json={"months": [{"month": "2025-04", "revenue": None}]}).status_code == 204
    months = client.get(url, params={"year": 2025}).json()
    assert months[0]["month"] == "2025-02" and len(months) == 11
    assert [(m["revenue"], m["earned"]) for m in months[:3]] == [(1000, 200), (0, 0), (None, 0)]
    assert months[-1]["revenue"] is None  # December is after the end date

    listed = next(s for s in client.get(f"/api/companies/{cid}/revenue-shares", params={"year": 2025}).json()
                  if s["id"] == sid)
    assert listed["earned_year"] == 200 and listed["missing_months"][0] == "2025-04"

    # Only reported-basis agreements take entered revenue.
    income_url = f"/api/companies/{cid}/revenue-shares/{investor['id']}/months"
    assert client.put(income_url, json={"months": [{"month": "2025-02", "revenue": 1}]}).status_code == 400


def test_validation_and_deletes(client):
    cid, app_a, _, investor, consultant, add = setup(client)
    url = f"/api/companies/{cid}/revenue-shares"
    base = {"payee": "P", "rate_percent": 5, "start_date": "2025-06-01"}
    assert client.post(url, json={**base, "end_date": "2025-05-31"}).status_code == 422
    assert client.post(url, json={**base, "rate_percent": 0}).status_code == 422
    assert client.post(url, json={**base, "payee": "  "}).status_code == 422
    other = client.post("/api/companies", json={"name": "Other"}).json()["id"]
    other_product = client.post(f"/api/companies/{other}/products", json={"name": "X"}).json()["id"]
    assert client.post(url, json={**base, "product_id": other_product}).status_code == 404

    # Only an expense can be a revenue-share payment.
    income = add(kind="income", date="2025-03-01", amount=5, revenue_share_id=investor["id"])
    assert income["revenue_share_id"] is None

    assert client.delete(f"{url}/{investor['id']}").status_code == 409
    assert client.delete(f"{url}/{consultant['id']}").status_code == 204

    # Deleting the company removes products, agreements, and payments with it.
    assert client.delete(f"/api/companies/{cid}").status_code == 204
    assert client.get(f"/api/companies/{cid}").status_code == 404


def test_statements_include_cost_of_revenue_products_and_shares(client):
    cid, *_ = setup(client)
    with SessionLocal() as db:
        st = build_statements(db, db.get(Company, cid), 2025, today=date(2026, 2, 1))
    assert [(l.label, l.amount) for l in st.cost_of_revenue] == [("App Store fees", 150), ("Revenue share", 85)]
    assert [(l.label, l.amount) for l in st.expenses] == [("Hosting", 40)]
    assert (st.total_income, st.gross_profit, st.net) == (D("3800.55"), D("3565.55"), D("3525.55"))
    assert [(p.name, p.net) for p in st.products] == [("App A", D("2765.55")), ("App B", 460), ("No product", 300)]
    assert [(s.payee, s.earned, s.paid, s.balance) for s in st.shares] == [
        ("Consultant", 0, 0, 0), ("Investor", D("285.06"), 85, D("200.06"))]
    assert st.shares[1].terms == "10% of App A revenue after App Store fees, starting January 1, 2025"
    assert st.ledger[-1].product == "App B"

    assert client.get(f"/api/companies/{cid}/statements", params={"year": 2025}).status_code == 200
    xlsx = client.get(f"/api/companies/{cid}/statements", params={"year": 2025, "format": "xlsx"})
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Income statement", "By product", "Members' capital", "Revenue shares",
                             "Cash summary", "Transactions"]
    labels = [c.value for c in wb["Income statement"]["A"]]
    assert "Gross profit" in labels and "Total operating expenses" in labels


def test_payee_statement(client):
    cid, app_a, _, investor, consultant, add = setup(client)
    add(kind="expense", date="2025-03-05", amount=150, category="Revenue share", revenue_share_id=investor["id"])
    with SessionLocal() as db:
        st = revenue_shares.build_payee_statement(
            db, db.get(Company, cid), db.get(RevenueShare, investor["id"]), date(2025, 3, 1), today=date(2025, 3, 20))
    assert st.period_label == "March 2025, through March 20, 2025"
    assert (st.opening, st.current.earned, st.current.balance) == (D("200.06"), 0, D("50.06"))
    assert [p.amount for p in st.month_payments] == [150]
    assert [r.month.month for r in st.year_to_date] == [1, 2, 3]

    url = f"/api/companies/{cid}/revenue-shares/{investor['id']}/statement"
    pdf = client.get(url, params={"month": "2025-02"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert 'filename="Apps LLC Investor revenue share 2025-02.pdf"' in pdf.headers["content-disposition"]
    xlsx = client.get(url, params={"month": "2025-02", "format": "xlsx"})
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Statement", "Year to date"]
    assert wb["Statement"]["B7"].value == 2000.55

    reported = f"/api/companies/{cid}/revenue-shares/{consultant['id']}/statement"
    assert client.get(reported, params={"month": "2025-01"}).status_code == 400  # before it starts
    assert client.get(reported, params={"month": "2025-03", "format": "xlsx"}).status_code == 200  # nothing entered
    assert client.get(url, params={"month": "2999-01"}).status_code == 400
    assert client.get(url, params={"month": "March"}).status_code == 422


def test_funding_is_a_liability_repaid_by_payments(client):
    """The investor funds the company; payments to them repay that first, then become expenses."""
    cid, app_a, _, investor, consultant, add = setup(client)
    sid = investor["id"]
    url = f"/api/companies/{cid}/entries"
    no_share = client.post(url, json={"kind": "funding", "date": "2025-01-05", "description": "x", "amount": 1})
    assert no_share.status_code == 400

    funded = add(kind="funding", date="2025-01-05", amount=100, revenue_share_id=sid, category="Sales",
                 product_id=app_a)
    assert (funded["category"], funded["product_id"], funded["revenue_share_id"]) == (None, None, sid)
    march = add(kind="expense", date="2025-03-05", amount=150, product_id=app_a, category="Revenue share",
                revenue_share_id=sid)
    add(kind="funding", date="2025-06-01", amount=50, revenue_share_id=sid)

    # Feb's $85 payment repays $85 of the $100; March's $150 repays the last $15, and $135 is an expense.
    repaid = {e["id"]: e["funding_repaid"] for e in client.get(url, params={"year": 2025}).json()}
    assert repaid[march["id"]] == 15 and sorted(v for v in repaid.values() if v) == [15, 85]

    # Funding isn't income, so it doesn't change what the investor earns.
    jan = client.get(f"/api/companies/{cid}/revenue-shares/{sid}/months", params={"year": 2025}).json()[0]
    assert (jan["revenue"], jan["earned"]) == (1000, 85)

    s = client.get(f"/api/companies/{cid}/summary", params={"year": 2025}).json()
    assert (s["income"], s["expenses"], s["net"]) == (3800.55, 325, 3475.55)  # 425 paid out, 100 repaid funding
    assert (s["funding_received"], s["funding_repaid"], s["cost_of_revenue"]) == (150, 100, 285)
    assert sum(m["expenses"] for m in s["months"]) == 325
    assert next(p for p in s["products"] if p["name"] == "App A")["cost_of_revenue"] == 285
    listed = next(c for c in client.get("/api/companies", params={"year": 2025}).json() if c["id"] == cid)
    assert listed["expenses"] == 325

    share = next(x for x in client.get(f"/api/companies/{cid}/revenue-shares", params={"year": 2025}).json()
                 if x["id"] == sid)
    assert (share["funding_received"], share["funding_repaid"], share["funding_outstanding"]) == (150, 100, 50)
    assert share["paid_total"] == 235  # payments are still payments to the investor

    with SessionLocal() as db:
        company = db.get(Company, cid)
        st = build_statements(db, company, 2025, today=date(2026, 2, 1))
        next_year = build_statements(db, company, 2026, today=date(2026, 2, 1))
        june = revenue_shares.build_payee_statement(db, company, db.get(RevenueShare, sid), date(2025, 6, 1),
                                                    today=date(2025, 7, 10))
        mar = revenue_shares.build_payee_statement(db, company, db.get(RevenueShare, sid), date(2025, 3, 1),
                                                   today=date(2025, 7, 10))
        consultant_st = revenue_shares.build_payee_statement(
            db, company, db.get(RevenueShare, consultant["id"]), date(2025, 3, 1), today=date(2025, 7, 10))
    assert [(f.payee, f.beginning, f.received, f.repaid, f.ending) for f in st.funding] == [
        ("Investor", 0, 150, 100, 50)]
    assert [(l.label, l.amount) for l in st.cost_of_revenue] == [("App Store fees", 150), ("Revenue share", 135)]
    assert st.net == D("3475.55") and st.net_cash_change == D("3525.55")
    assert sum(r.amount for r in st.ledger) == st.net_cash_change
    assert [(f.beginning, f.ending) for f in next_year.funding] == [(50, 50)]
    assert (june.funding.beginning, june.funding.received, june.funding.ending) == (0, 50, 50)
    assert (mar.funding.beginning, mar.funding.repaid, mar.funding.ending) == (15, 15, 0)
    assert consultant_st.funding is None

    xlsx = client.get(f"/api/companies/{cid}/statements", params={"year": 2025, "format": "xlsx"})
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert "Funding" in wb.sheetnames
    assert "Revenue-share funding repaid" in [c.value for c in wb["Cash summary"]["A"]]
    assert client.get(f"/api/companies/{cid}/statements", params={"year": 2025}).status_code == 200
    statement = f"/api/companies/{cid}/revenue-shares/{sid}/statement"
    assert client.get(statement, params={"month": "2025-06"}).status_code == 200
    assert client.get(statement, params={"month": "2025-06", "format": "xlsx"}).status_code == 200

    # An agreement with funding recorded can't be deleted.
    add(kind="funding", date="2025-03-01", amount=5, revenue_share_id=consultant["id"])
    assert client.delete(f"/api/companies/{cid}/revenue-shares/{consultant['id']}").status_code == 409
