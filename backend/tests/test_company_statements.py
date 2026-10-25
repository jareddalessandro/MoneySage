import io
from datetime import date

from openpyxl import load_workbook

from app.database import SessionLocal
from app.models import Company
from app.services.company_statements import build_statements


def setup_company(client):
    cid = client.post("/api/companies", json={"name": "Acme / Co"}).json()["id"]
    me = client.post(f"/api/companies/{cid}/members", json={"name": "Jared", "equity_percent": 60}).json()
    sam = client.post(f"/api/companies/{cid}/members", json={"name": "Sam", "equity_percent": 30}).json()
    entries = [
        # 2025: the prior year, which makes up the beginning balances.
        dict(kind="contribution", date="2025-01-10", amount=5000, member_id=me["id"]),
        dict(kind="income", date="2025-06-01", amount=2000, category="Consulting"),
        dict(kind="expense", date="2025-07-01", amount=1000, category="Software"),
        dict(kind="distribution", date="2025-12-20", amount=200, member_id=sam["id"]),
        # 2026
        dict(kind="income", date="2026-02-01", amount=8000, category="Consulting"),
        dict(kind="income", date="2026-03-01", amount=500),
        dict(kind="expense", date="2026-03-15", amount=1200, category="Software"),
        dict(kind="expense", date="2026-04-15", amount=300, category="Fees"),
        dict(kind="contribution", date="2026-01-05", amount=1000, member_id=sam["id"]),
        dict(kind="distribution", date="2026-06-30", amount=2000, member_id=me["id"]),
        # 2027: after the statement year, ignored.
        dict(kind="income", date="2027-01-02", amount=999),
    ]
    for e in entries:
        assert client.post(f"/api/companies/{cid}/entries", json={"description": "x", **e}).status_code == 201
    return cid


def test_statement_figures(client):
    cid = setup_company(client)
    with SessionLocal() as db:
        st = build_statements(db, db.get(Company, cid), 2026, today=date(2027, 2, 1))

    assert st.is_full_year and st.period_label == "For the year ended December 31, 2026"
    assert [(l.label, l.amount) for l in st.income] == [("Consulting", 8000), ("Uncategorized", 500)]
    assert [(l.label, l.amount) for l in st.expenses] == [("Software", 1200), ("Fees", 300)]
    assert st.net == 7000

    jared, sam = st.capital
    # Beginning = prior contributions - prior distributions + prior net ($1,000) x ownership.
    assert (jared.beginning, sam.beginning) == (5600, 100)
    assert (jared.contributions, jared.profit_share, jared.distributions, jared.ending) == (0, 4200, 2000, 7800)
    assert (sam.contributions, sam.profit_share, sam.distributions, sam.ending) == (1000, 2100, 0, 3200)
    assert st.unallocated_profit == 700  # ownership totals 90%

    # Cash: prior = 5000 + 2000 - 1000 - 200; this year = 8500 + 1000 - 1500 - 2000.
    assert (st.cash_beginning, st.net_cash_change, st.cash_ending) == (5800, 6000, 11800)
    assert len(st.ledger) == 6 and st.ledger[0].date == date(2026, 1, 5)
    assert [r.amount for r in st.ledger if r.kind == "Distribution"] == [-2000]

    with SessionLocal() as db:
        ytd = build_statements(db, db.get(Company, cid), 2026, today=date(2026, 10, 24))
    assert not ytd.is_full_year and ytd.period_label == "Year to date, January 1 through October 24, 2026"


def test_statement_downloads(client):
    cid = setup_company(client)

    pdf = client.get(f"/api/companies/{cid}/statements", params={"year": 2025, "format": "pdf"})
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")
    assert 'filename="Acme  Co 2025 financial statements.pdf"' in pdf.headers["content-disposition"]

    xlsx = client.get(f"/api/companies/{cid}/statements", params={"year": 2025, "format": "xlsx"})
    assert xlsx.status_code == 200
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Income statement", "Members' capital", "Cash summary", "Transactions"]
    capital = wb["Members' capital"]
    assert [c.value for c in capital[7]][:7] == ["Jared", 0.6, 0, 5000, 600, 0, 5600]
    assert capital["G9"].value == "=SUM(G7:G8)"
    transactions = wb["Transactions"]
    assert [transactions.cell(row=r, column=6).value for r in range(7, 11)] == [5000, 2000, -1000, -200]

    # Text with markup characters doesn't break the PDF, and an empty company still exports.
    client.post("/api/companies", json={"name": "<b>Empty</b> & Co"})
    empty_id = client.get("/api/companies").json()[0]["id"]
    assert client.get(f"/api/companies/{empty_id}/statements").status_code == 200
    assert client.get(f"/api/companies/{empty_id}/statements", params={"format": "xlsx"}).status_code == 200
    assert client.get(f"/api/companies/{cid}/statements", params={"format": "csv"}).status_code == 422
