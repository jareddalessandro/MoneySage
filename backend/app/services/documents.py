"""PDF (reportlab) and Excel (openpyxl) helpers shared by the company's downloadable documents."""

import io
from decimal import Decimal
from xml.sax.saxutils import escape


def money(n: Decimal) -> str:
    """Accounting style: $1,234.56, and ($1,234.56) for negatives."""
    text = f"${abs(n):,.2f}"
    return f"({text})" if n < 0 else text


def pct(n: Decimal) -> str:
    return f"{n.normalize():f}%"


class PdfKit:
    """Letter-size page, styles, and statement tables."""

    def __init__(self):
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_RIGHT
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch

        self.ink = colors.HexColor("#1b1f24")
        self.muted = colors.HexColor("#5f6b73")
        self.band = colors.HexColor("#eef1f4")
        self.rule = colors.HexColor("#c9d0d6")
        base = getSampleStyleSheet()
        self.title = ParagraphStyle("title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=20,
                                    leading=24, alignment=0, textColor=self.ink, spaceAfter=2)
        self.subtitle = ParagraphStyle("subtitle", parent=base["Normal"], fontSize=11, leading=15, textColor=self.ink)
        self.small = ParagraphStyle("small", parent=base["Normal"], fontSize=8.5, leading=11, textColor=self.muted)
        self.heading = ParagraphStyle("heading", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=13,
                                      leading=16, textColor=self.ink, spaceBefore=16, spaceAfter=6)
        self.cell = ParagraphStyle("cell", parent=base["Normal"], fontSize=9, leading=11, textColor=self.ink)
        self.cell_right = ParagraphStyle("cell_right", parent=self.cell, alignment=TA_RIGHT)
        self.page = LETTER
        self.inch = inch
        self.width = LETTER[0] - 1.5 * inch

    def p(self, text: str, style=None):
        """A paragraph of plain text. Escaped, because Paragraph parses markup."""
        from reportlab.platypus import Paragraph

        return Paragraph(escape(text), style or self.cell)

    def table(self, rows, col_widths, totals=(), sections=(), grand=(), indent=()):
        """rows: list of lists. totals/sections/grand/indent: row indexes to style."""
        from reportlab.platypus import Table, TableStyle

        t = Table(rows, colWidths=col_widths, repeatRows=1)
        style = [
            ("FONT", (0, 0), (-1, -1), "Helvetica", 9.5),
            ("TEXTCOLOR", (0, 0), (-1, -1), self.ink),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
            ("TEXTCOLOR", (0, 0), (-1, 0), self.muted),
            ("LINEBELOW", (0, 0), (-1, 0), 0.75, self.rule),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
        for i in sections:
            style += [("FONT", (0, i), (-1, i), "Helvetica-Bold", 9.5), ("BACKGROUND", (0, i), (-1, i), self.band)]
        for i in totals:
            style += [("FONT", (0, i), (-1, i), "Helvetica-Bold", 9.5), ("LINEABOVE", (1, i), (-1, i), 0.5, self.rule)]
        for i in grand:
            style += [
                ("FONT", (0, i), (-1, i), "Helvetica-Bold", 10.5),
                ("LINEABOVE", (0, i), (-1, i), 0.75, self.ink),
                ("LINEBELOW", (0, i), (-1, i), 1.5, self.ink),
            ]
        for i in indent:
            style.append(("LEFTPADDING", (0, i), (0, i), 18))
        t.setStyle(TableStyle(style))
        return t

    def build(self, story, footer_text: str, title: str) -> bytes:
        from reportlab.platypus import SimpleDocTemplate

        inch, page = self.inch, self.page

        def footer(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(self.muted)
            canvas.drawString(0.75 * inch, 0.5 * inch, footer_text)
            canvas.drawRightString(page[0] - 0.75 * inch, 0.5 * inch, f"Page {doc.page}")
            canvas.restoreState()

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=page, leftMargin=0.75 * inch, rightMargin=0.75 * inch, topMargin=0.75 * inch,
            bottomMargin=0.8 * inch, title=title, author="MoneySage",
        )
        doc.build(story, onFirstPage=footer, onLaterPages=footer)
        return buffer.getvalue()


class XlsxKit:
    """A workbook whose sheets share four header lines (name, title, period, prepared note)."""

    money_fmt = '"$"#,##0.00_);("$"#,##0.00)'
    pct_fmt = "0.00##%"

    def __init__(self, heading: str, period: str, prepared: str):
        from openpyxl import Workbook
        from openpyxl.styles import Border, Font, PatternFill, Side

        self.wb = Workbook()
        self.heading, self.period, self.prepared = heading, period, prepared
        self.bold = Font(bold=True)
        self.header_font = Font(bold=True, color="5F6B73", size=9)
        self.note_font = Font(italic=True, color="5F6B73", size=9)
        self.section_fill = PatternFill("solid", fgColor="EEF1F4")
        self.top_rule = Border(top=Side(style="thin"))
        self.grand_rule = Border(top=Side(style="thin"), bottom=Side(style="double"))

    def sheet(self, name: str, title: str, widths: list[int]):
        """A sheet with the header lines; content starts on row 6."""
        from openpyxl.styles import Font

        ws = self.wb.active if self.wb.active.title == "Sheet" else self.wb.create_sheet()
        ws.title = name
        ws["A1"] = self.heading
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = title
        ws["A2"].font = self.bold
        ws["A3"] = self.period
        ws["A4"] = self.prepared
        ws["A4"].font = self.note_font
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = w
        return ws

    def put(self, ws, row: int, values: list, *, font=None, fill=None, border=None, money_cols=(), pct_cols=(),
            indent=False):
        from openpyxl.styles import Alignment

        for col, value in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=value)
            if col in money_cols:
                c.number_format = self.money_fmt
            if col in pct_cols:
                c.number_format = self.pct_fmt
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if border and col > 1:
                c.border = border
        if indent:
            ws.cell(row=row, column=1).alignment = Alignment(indent=1)

    def note(self, ws, row: int, text: str):
        ws.cell(row=row, column=1, value=text).font = self.note_font

    def save(self) -> bytes:
        buffer = io.BytesIO()
        self.wb.save(buffer)
        return buffer.getvalue()
